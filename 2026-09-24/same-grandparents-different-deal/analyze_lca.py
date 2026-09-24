#!/usr/bin/env python3
"""Stream LCA FY2025 Q4 xlsx -> median offered wages for Big Tech software roles
in Bay Area worksites (certified cases)."""
import json
import re
from pathlib import Path

import openpyxl

PROJ = Path("/home/hatch/workspace/passion-projects/2026-09-24/same-grandparents-different-deal")
RAW = PROJ / "data" / "raw"

EMPLOYERS = {
    "META": "Meta", "GOOGLE": "Google", "APPLE": "Apple", "AMAZON": "Amazon",
    "MICROSOFT": "Microsoft", "NVIDIA": "Nvidia", "NETFLIX": "Netflix",
    "ADOBE": "Adobe", "SALESFORCE": "Salesforce", "INTEL": "Intel",
    "ORACLE": "Oracle", "CISCO": "Cisco", "UBER": "Uber", "AIRBNB": "Airbnb",
    "STRIPE": "Stripe", "DOORDASH": "DoorDash", "LYFT": "Lyft",
}
BAY = {"SAN FRANCISCO", "SAN JOSE", "MOUNTAIN VIEW", "MENLO PARK", "CUPERTINO",
       "SUNNYVALE", "PALO ALTO", "SANTA CLARA", "REDWOOD CITY", "FREMONT",
       "OAKLAND", "BERKELEY", "SAN MATEO", "BELMONT", "LOS ALTOS"}

MULT = {"HOUR": 2080, "WEEK": 52, "BI-WEEKLY": 26, "MONTH": 12, "YEAR": 1,
        "2 WEEKS": 26}


def classify_emp(name):
    n = (name or "").upper()
    for key, label in EMPLOYERS.items():
        if key in n:
            return label
    return None


def annualize(val, unit):
    try:
        v = float(val)
    except (TypeError, ValueError):
        return None
    m = MULT.get((unit or "").upper().strip())
    return v * m if m else None


def main():
    wb = openpyxl.load_workbook(RAW / "lca_fy2025_q4.xlsx", read_only=True, data_only=True)
    ws = wb.active
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    I = {h: i for i, h in enumerate(hdr)}
    by_emp = {}
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        n += 1
        if str(row[I["CASE_STATUS"]] or "").strip().upper() != "CERTIFIED":
            continue
        emp = classify_emp(row[I["EMPLOYER_NAME"]])
        if not emp:
            continue
        if str(row[I["WORKSITE_STATE"]] or "").upper().strip() not in ("CA", "CALIFORNIA"):
            continue
        if str(row[I["WORKSITE_CITY"]] or "").upper().strip() not in BAY:
            continue
        title = str(row[I["JOB_TITLE"]] or "").upper()
        if "SOFTWARE" not in title:
            continue
        w = annualize(row[I["WAGE_RATE_OF_PAY_FROM"]], row[I["WAGE_UNIT_OF_PAY"]])
        if w and 50000 < w < 1000000:
            by_emp.setdefault(emp, []).append(w)
    wb.close()
    print(f"scanned {n} rows")
    summary = {}
    for emp, wages in sorted(by_emp.items(), key=lambda x: -len(x[1])):
        wages.sort()
        med = wages[len(wages) // 2]
        summary[emp] = {"n": len(wages), "median": round(med)}
        print(f"{emp}: n={len(wages)} median=${med:,.0f}")
    json.dump(summary, open(PROJ / "data" / "lca_tech_wages.json", "w"), indent=2)


if __name__ == "__main__":
    main()
